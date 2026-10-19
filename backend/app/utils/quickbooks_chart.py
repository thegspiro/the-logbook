"""Reading a QuickBooks Online chart of accounts out of an exported report.

The Treasurer exports the **Account List** report from QuickBooks Online
(Reports › Account List › Export, to Excel or CSV) and uploads it, so the
QuickBooks export can check its account names against the accounts that
actually exist before a file is imported.

The report is not a bare table: it opens with a title block (the company
name, "Account List") and can close with a footer line (when it was run).
So the header row is found by its column names rather than assumed to be
the first, and a row with no account type is taken for title or footer text
rather than an account. Intuit has labelled the name column both "Account"
and "Full name"; both are read, and the full name (subaccounts written
Parent:Child, the form the journal-entry import expects) wins when a file
carries both.
"""

import csv
import io
import zipfile
from dataclasses import dataclass
from typing import Iterable, Iterator, Optional, Sequence

# The upload is a report, not a ledger: a department with a few hundred
# accounts produces a file of tens of kilobytes. The caps bound a hostile or
# mistaken upload rather than shaping a real one.
MAX_CHART_FILE_BYTES = 5 * 1024 * 1024
MAX_CHART_ACCOUNTS = 5000
# A report row beyond this is not part of any chart of accounts.
_MAX_ROWS = 20_000
# The title block above the header is a handful of rows.
_HEADER_SEARCH_ROWS = 20
# The longest name the model stores.
MAX_ACCOUNT_NAME_LENGTH = 255

# Header labels naming the account, best first.
_NAME_HEADERS = ("full name", "account full name", "account name", "account", "name")
_TYPE_HEADERS = ("type", "account type")


@dataclass(frozen=True)
class ChartAccount:
    name: str
    account_type: Optional[str]


def _label(cell: object) -> str:
    return " ".join(str(cell or "").split()).casefold()


def _text(cell: object) -> str:
    return str(cell).strip() if cell is not None else ""


def _csv_rows(raw: bytes) -> Iterator[Sequence[object]]:
    try:
        content = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        content = raw.decode("latin-1")
    yield from csv.reader(io.StringIO(content))


def _xlsx_rows(raw: bytes) -> Iterator[Sequence[object]]:
    # openpyxl parses the workbook's XML through defusedxml when it is
    # installed, which requirements.txt pins for exactly this reason.
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, OSError, ValueError) as exc:
        raise ValueError("The file could not be read as an Excel workbook.") from exc
    try:
        sheet = workbook.worksheets[0] if workbook.worksheets else None
        if sheet is None:
            return
        for row in sheet.iter_rows(values_only=True):
            yield row
    finally:
        workbook.close()


def parse_chart_of_accounts(raw: bytes, filename: str) -> list[ChartAccount]:
    """The accounts in an exported Account List, in file order, de-duplicated.

    Raises ``ValueError`` with a message fit for the Treasurer when the file
    is not a readable report or holds no accounts.
    """
    lower = (filename or "").lower()
    rows: Iterable[Sequence[object]]
    if lower.endswith(".xlsx"):
        rows = _xlsx_rows(raw)
    elif lower.endswith(".csv"):
        rows = _csv_rows(raw)
    else:
        raise ValueError("Upload the Account List report as an .xlsx or .csv file.")

    name_col: Optional[int] = None
    type_col: Optional[int] = None
    accounts: list[ChartAccount] = []
    seen: set[str] = set()
    for index, row in enumerate(rows):
        if index >= _MAX_ROWS:
            raise ValueError(
                f"The file has more than {_MAX_ROWS:,} rows; it does not look "
                "like an Account List report."
            )
        if name_col is None:
            if index >= _HEADER_SEARCH_ROWS:
                break
            labels = [_label(cell) for cell in row]
            for wanted in _NAME_HEADERS:
                if wanted in labels:
                    name_col = labels.index(wanted)
                    break
            if name_col is not None:
                type_col = next(
                    (labels.index(t) for t in _TYPE_HEADERS if t in labels), None
                )
            continue

        name = _text(row[name_col]) if name_col < len(row) else ""
        account_type = (
            _text(row[type_col]) if type_col is not None and type_col < len(row) else ""
        )
        # Title, total and footer lines carry text in the first column and
        # nothing in Type; every account in the report has a type.
        if not name or (type_col is not None and not account_type):
            continue
        if len(name) > MAX_ACCOUNT_NAME_LENGTH:
            raise ValueError(
                f"An account name in the file is longer than "
                f"{MAX_ACCOUNT_NAME_LENGTH} characters."
            )
        if name in seen:
            continue
        seen.add(name)
        accounts.append(ChartAccount(name, account_type[:100] or None))
        if len(accounts) > MAX_CHART_ACCOUNTS:
            raise ValueError(
                f"The file lists more than {MAX_CHART_ACCOUNTS:,} accounts; it "
                "does not look like an Account List report."
            )

    if name_col is None:
        raise ValueError(
            "No account column was found. Upload the Account List report from "
            "QuickBooks Online; its header row names an Account or Full name "
            "column."
        )
    if not accounts:
        raise ValueError("The file lists no accounts.")
    return accounts
